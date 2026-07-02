#!/usr/bin/env python3
"""Fetch latest nginx/redis/nsq versions and update daily Excel report."""

from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from packaging import version

HEADERS = ["日期", "nginx", "redis", "nsq", "是否存在漏洞", "修复方法"]
EXCEL_NAME = "中间件版本信息.xlsx"
REQUEST_TIMEOUT = 30
USER_AGENT = "middleware-version-tracker/1.0"


@dataclass
class MiddlewareInfo:
    name: str
    version: str
    has_vulnerability: bool
    fix_methods: list[str]


def get_output_dirs() -> list[Path]:
    candidates = [
        Path.home() / "Desktop" / "ztx",
        Path.home() / "桌面" / "ztx",
        Path(__file__).resolve().parent.parent / "ztx",
    ]
    output_dirs: list[Path] = []
    seen: set[str] = set()
    for path in candidates:
        resolved = str(path.resolve())
        if resolved in seen:
            continue
        seen.add(resolved)
        path.mkdir(parents=True, exist_ok=True)
        output_dirs.append(path)
    if not output_dirs:
        raise RuntimeError("Unable to determine output directory")
    return output_dirs


def fetch_nginx_version() -> str:
    response = requests.get(
        "https://nginx.org/en/download.html",
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT},
    )
    response.raise_for_status()
    stable_section = response.text.split("Stable version", 1)
    search_text = stable_section[1] if len(stable_section) > 1 else response.text
    match = re.search(r"nginx-(\d+\.\d+\.\d+)", search_text)
    if not match:
        match = re.search(r"nginx-(\d+\.\d+\.\d+)", response.text)
    if not match:
        raise RuntimeError("Failed to parse nginx stable version")
    return match.group(1)


def fetch_github_latest_version(repo: str) -> str:
    response = requests.get(
        f"https://api.github.com/repos/{repo}/releases/latest",
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT, "Accept": "application/vnd.github+json"},
    )
    response.raise_for_status()
    tag = response.json()["tag_name"]
    return tag.lstrip("v")


def parse_nginx_not_vulnerable() -> list[str]:
    """Return nginx advisory 'Not vulnerable' version ranges."""
    response = requests.get(
        "https://nginx.org/en/security_advisories.html",
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT},
    )
    response.raise_for_status()
    pattern = re.compile(
        r"Not vulnerable:\s*([^V]+?)Vulnerable:",
        re.IGNORECASE,
    )
    ranges: list[str] = []
    for match in pattern.finditer(response.text):
        ranges.extend(part.strip() for part in match.group(1).split(",") if part.strip())
    return ranges


def version_matches_not_vulnerable(current: str, not_vulnerable_ranges: Iterable[str]) -> bool:
    current_v = version.parse(current)
    for item in not_vulnerable_ranges:
        item = item.strip()
        if item.endswith("+"):
            base = item[:-1].strip()
            if current_v >= version.parse(base):
                return True
            continue
        if "-" in item:
            start, end = item.split("-", 1)
            if version.parse(start) <= current_v <= version.parse(end):
                return True
    return False


def check_nginx_vulnerabilities(nginx_version: str) -> tuple[bool, list[str]]:
    not_vulnerable = parse_nginx_not_vulnerable()
    if version_matches_not_vulnerable(nginx_version, not_vulnerable):
        return False, []
    return True, [
        f"升级 nginx 至最新稳定版（当前最新: {nginx_version}），参考 https://nginx.org/en/security_advisories.html"
    ]


def fetch_redis_latest_release() -> tuple[str, str]:
    response = requests.get(
        "https://api.github.com/repos/redis/redis/releases/latest",
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT, "Accept": "application/vnd.github+json"},
    )
    response.raise_for_status()
    payload = response.json()
    return payload["tag_name"].lstrip("v"), payload["published_at"]


def fetch_redis_fixed_cves(redis_version: str) -> set[str]:
    major_minor = ".".join(redis_version.split(".")[:2])
    urls = [
        f"https://raw.githubusercontent.com/redis/redis/{redis_version}/00-RELEASENOTES",
        f"https://raw.githubusercontent.com/redis/redis/{major_minor}/00-RELEASENOTES",
    ]
    fixed: set[str] = set()
    for url in urls:
        try:
            response = requests.get(
                url,
                timeout=REQUEST_TIMEOUT,
                headers={"User-Agent": USER_AGENT},
            )
            if response.status_code != 200:
                continue
            fixed.update(re.findall(r"CVE-\d{4}-\d+", response.text))
        except requests.RequestException:
            continue
    return fixed


def fetch_redis_advisories() -> list[dict]:
    response = requests.get(
        "https://api.github.com/repos/redis/redis/security-advisories",
        timeout=REQUEST_TIMEOUT,
        headers={"User-Agent": USER_AGENT, "Accept": "application/vnd.github+json"},
    )
    response.raise_for_status()
    return response.json()


def check_redis_vulnerabilities(redis_version: str) -> tuple[bool, list[str]]:
    current = version.parse(redis_version)
    latest_version, latest_published_at = fetch_redis_latest_release()
    on_latest_release = redis_version == latest_version
    fixed_cves = fetch_redis_fixed_cves(redis_version)
    fixes: list[str] = []
    for advisory in fetch_redis_advisories():
        patched = advisory.get("patched_versions") or []
        affected = advisory.get("vulnerabilities") or []
        cve_id = (advisory.get("cve_id") or advisory.get("ghsa_id") or "未知CVE").strip()
        summary = (advisory.get("summary") or "").strip()
        published_at = advisory.get("published_at") or ""

        if cve_id in fixed_cves:
            continue
        if on_latest_release and published_at and published_at <= latest_published_at:
            continue

        is_patched = any(
            _redis_version_satisfies(current, patched_version)
            for patched_version in patched
        )
        if is_patched:
            continue

        for item in affected:
            item_patched = item.get("patched_versions")
            if isinstance(item_patched, str) and item_patched not in {"", "TBD"}:
                if _redis_version_satisfies(current, item_patched):
                    is_patched = True
                    break
        if is_patched:
            continue

        affects_current = any(
            _redis_version_in_range(current, item.get("vulnerable_version_range", ""))
            for item in affected
        )
        if affects_current:
            patched_hint = ", ".join(patched) if patched else "最新版"
            fix_hint = f"升级 redis 至 {patched_hint} 修复 {cve_id}"
            if summary:
                fix_hint += f"（{summary}）"
            fixes.append(fix_hint)

    return bool(fixes), fixes


def _safe_parse_version(value: str) -> version.Version | None:
    value = value.strip()
    if not value or value.lower() == "all":
        return None
    try:
        return version.parse(value)
    except version.InvalidVersion:
        return None


def _redis_version_satisfies(current: version.Version, patched_version: str) -> bool:
    patched_version = patched_version.strip()
    if patched_version.startswith(">="):
        target = _safe_parse_version(patched_version[2:])
        return target is not None and current >= target
    if patched_version.startswith("<="):
        target = _safe_parse_version(patched_version[2:])
        return target is not None and current <= target
    if patched_version.startswith(">"):
        target = _safe_parse_version(patched_version[1:])
        return target is not None and current > target
    if patched_version.startswith("<"):
        target = _safe_parse_version(patched_version[1:])
        return target is not None and current < target
    target = _safe_parse_version(patched_version)
    return target is not None and current == target


def _redis_version_in_range(current: version.Version, vulnerable_range: str) -> bool:
    vulnerable_range = vulnerable_range.strip()
    if not vulnerable_range:
        return False
    if vulnerable_range.lower() == "all":
        return True
    if vulnerable_range.startswith("<"):
        target = _safe_parse_version(vulnerable_range[1:])
        return target is not None and current < target
    if vulnerable_range.startswith("<="):
        target = _safe_parse_version(vulnerable_range[2:])
        return target is not None and current <= target
    if vulnerable_range.startswith(">="):
        target = _safe_parse_version(vulnerable_range[2:])
        return target is not None and current >= target
    if vulnerable_range.startswith(">"):
        target = _safe_parse_version(vulnerable_range[1:])
        return target is not None and current > target
    if " - " in vulnerable_range:
        start, end = vulnerable_range.split(" - ", 1)
        start_v = _safe_parse_version(start)
        end_v = _safe_parse_version(end)
        if start_v is None or end_v is None:
            return False
        return start_v <= current <= end_v
    target = _safe_parse_version(vulnerable_range)
    return target is not None and current == target


def check_nsq_vulnerabilities(nsq_version: str) -> tuple[bool, list[str]]:
    try:
        response = requests.get(
            "https://api.github.com/repos/nsqio/nsq/security-advisories",
            timeout=REQUEST_TIMEOUT,
            headers={"User-Agent": USER_AGENT, "Accept": "application/vnd.github+json"},
        )
        if response.status_code != 200:
            return False, []
        advisories = response.json()
    except requests.RequestException:
        return False, []

    fixes: list[str] = []
    current = version.parse(nsq_version)
    for advisory in advisories:
        patched = advisory.get("patched_versions") or []
        affected = advisory.get("vulnerabilities") or []
        cve_id = (advisory.get("cve_id") or advisory.get("ghsa_id") or "未知CVE").strip()
        summary = (advisory.get("summary") or "").strip()

        is_patched = any(
            _redis_version_satisfies(current, patched_version)
            for patched_version in patched
        )
        if is_patched:
            continue

        affects_current = any(
            _redis_version_in_range(current, item.get("vulnerable_version_range", ""))
            for item in affected
        )
        if affects_current:
            fix_hint = f"升级 nsq 至 {', '.join(patched) or '最新版'} 修复 {cve_id}"
            if summary:
                fix_hint += f"（{summary}）"
            fixes.append(fix_hint)

    return bool(fixes), fixes


def collect_middleware_info() -> list[MiddlewareInfo]:
    nginx_version = fetch_nginx_version()
    redis_version = fetch_github_latest_version("redis/redis")
    nsq_version = fetch_github_latest_version("nsqio/nsq")

    nginx_vuln, nginx_fixes = check_nginx_vulnerabilities(nginx_version)
    redis_vuln, redis_fixes = check_redis_vulnerabilities(redis_version)
    nsq_vuln, nsq_fixes = check_nsq_vulnerabilities(nsq_version)

    return [
        MiddlewareInfo("nginx", nginx_version, nginx_vuln, nginx_fixes),
        MiddlewareInfo("redis", redis_version, redis_vuln, redis_fixes),
        MiddlewareInfo("nsq", nsq_version, nsq_vuln, nsq_fixes),
    ]


def build_row(today: date, items: list[MiddlewareInfo]) -> list[str]:
    has_any_vuln = any(item.has_vulnerability for item in items)
    fix_methods: list[str] = []
    for item in items:
        if item.has_vulnerability:
            fix_methods.extend(item.fix_methods)

    return [
        today.isoformat(),
        next(item.version for item in items if item.name == "nginx"),
        next(item.version for item in items if item.name == "redis"),
        next(item.version for item in items if item.name == "nsq"),
        "是" if has_any_vuln else "否",
        "\n".join(fix_methods) if fix_methods else "无",
    ]


def init_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "中间件版本"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 14
    sheet.column_dimensions["F"].width = 80
    workbook.save(path)


def upsert_daily_row(path: Path, row: list[str]) -> None:
    if not path.exists():
        init_workbook(path)

    workbook = load_workbook(path)
    sheet = workbook.active

    if sheet.max_row == 0:
        sheet.append(HEADERS)

    today = row[0]
    target_row = None
    for row_idx in range(2, sheet.max_row + 1):
        if str(sheet.cell(row=row_idx, column=1).value) == today:
            target_row = row_idx
            break

    if target_row is None:
        sheet.append(row)
        target_row = sheet.max_row
    else:
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=target_row, column=col_idx, value=value)

    sheet.cell(row=target_row, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def main() -> int:
    output_dirs = get_output_dirs()
    today = date.today()
    items = collect_middleware_info()
    row = build_row(today, items)

    for output_dir in output_dirs:
        excel_path = output_dir / EXCEL_NAME
        upsert_daily_row(excel_path, row)
        print(f"Updated: {excel_path}")

    print(
        " | ".join(
            [
                f"date={row[0]}",
                f"nginx={row[1]}",
                f"redis={row[2]}",
                f"nsq={row[3]}",
                f"vuln={row[4]}",
            ]
        )
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
